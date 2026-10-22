import csv
import io
from uuid import UUID

import openpyxl
import pytest

from app.models import AuditEvent


def create_company(client, code: str, prefix: str = 'AST'):
    response = client.post('/api/v1/companies', json={'code': code, 'name': f'{code} Corp', 'number_prefix': prefix})
    assert response.status_code == 201, response.text
    return response.json()


def setup_company_infrastructure(client, company_id: str):
    dept_res = client.post(
        f'/api/v1/companies/{company_id}/departments',
        json={'code': 'IT', 'name': 'Information Technology'},
    )
    assert dept_res.status_code == 201, dept_res.text
    dept_id = dept_res.json()['id']

    site_res = client.post(
        f'/api/v1/companies/{company_id}/sites',
        json={'code': 'HQ', 'name': 'Headquarters', 'address': '100 Main St'},
    )
    assert site_res.status_code == 201, site_res.text
    site_id = site_res.json()['id']

    loc_res = client.post(
        f'/api/v1/companies/{company_id}/sites/{site_id}/locations',
        json={'code': 'ROOM-1', 'name': 'Server Room 1'},
    )
    assert loc_res.status_code == 201, loc_res.text
    loc_id = loc_res.json()['id']

    cc_res = client.post(
        f'/api/v1/companies/{company_id}/cost-centers',
        json={'code': 'CC-IT', 'name': 'IT Cost Center'},
    )
    assert cc_res.status_code == 201, cc_res.text
    cc_id = cc_res.json()['id']

    return {
        'dept_id': dept_id,
        'site_id': site_id,
        'loc_id': loc_id,
        'cc_id': cc_id,
    }


def test_download_import_template_csv(client):
    company = create_company(client, 'TMPL-CSV')
    response = client.get(f"/api/v1/companies/{company['id']}/assets/import/template?format=csv")
    assert response.status_code == 200
    assert 'text/csv' in response.headers['content-type']
    assert 'attachment; filename="mindzen_asset_import_template.csv"' in response.headers['content-disposition']

    content = response.content.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))
    assert len(reader) >= 3  # Header + 2 sample rows
    headers = reader[0]
    assert 'category_code' in headers
    assert 'serial_number' in headers
    assert 'manufacturer' in headers
    assert 'model' in headers
    assert 'site_code' in headers
    assert 'location_code' in headers


def test_download_import_template_xlsx(client):
    company = create_company(client, 'TMPL-XLSX')
    response = client.get(f"/api/v1/companies/{company['id']}/assets/import/template?format=xlsx")
    assert response.status_code == 200
    assert 'application/vnd.openxmlformats' in response.headers['content-type']
    assert 'mindzen_asset_import_template.xlsx' in response.headers['content-disposition']

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) >= 3
    headers = rows[0]
    assert 'category_code' in headers
    assert 'serial_number' in headers


def test_csv_bulk_import_preview_and_confirm(client, db_session):
    company = create_company(client, 'BULK-CSV', prefix='CSV')
    setup = setup_company_infrastructure(client, company['id'])

    # Build CSV file
    csv_data = [
        [
            'category_code', 'serial_number', 'manufacturer', 'model', 'status',
            'site_code', 'location_code', 'department_code', 'cost_center_code',
            'acquisition_date', 'acquisition_cost', 'custodian_name',
            'vendor_reference', 'invoice_reference', 'warranty_start', 'warranty_end', 'notes',
        ],
        [
            'LAPTOP', 'SN-CSV-001', 'Apple', 'MacBook Pro 16', 'in_stock',
            'HQ', 'ROOM-1', 'IT', 'CC-IT',
            '2026-02-01', '2499.00', 'Jane Doe',
            'Apple Inc', 'INV-9901', '2026-02-01', '2029-02-01', 'Developer laptop',
        ],
        [
            'DESKTOP', 'SN-CSV-002', 'Dell', 'OptiPlex 7090', 'in_stock',
            'HQ', 'ROOM-1', 'IT', 'CC-IT',
            '2026-02-01', '1099.50', 'John Smith',
            'Dell Direct', 'INV-9902', '2026-02-01', '2029-02-01', 'Workstation desktop',
        ],
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows(csv_data)
    csv_bytes = output.getvalue().encode('utf-8')

    # Step 1: Preview
    preview_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('assets_batch.csv', csv_bytes, 'text/csv')},
    )
    assert preview_res.status_code == 200, preview_res.text
    preview_data = preview_res.json()
    assert preview_data['valid'] is True
    assert preview_data['total_rows'] == 2
    assert preview_data['valid_rows'] == 2
    assert preview_data['invalid_rows'] == 0
    assert len(preview_data['errors']) == 0
    assert len(preview_data['preview_rows']) == 2
    preview_token = preview_data['preview_token']
    assert preview_token is not None

    # Step 2: Confirm
    confirm_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/confirm",
        json={'preview_token': preview_token},
    )
    assert confirm_res.status_code == 200, confirm_res.text
    confirm_data = confirm_res.json()
    assert confirm_data['imported_count'] == 2
    assert len(confirm_data['asset_numbers']) == 2

    # Verify assets in registry
    assets_res = client.get(f"/api/v1/companies/{company['id']}/assets")
    assert assets_res.status_code == 200
    assets = assets_res.json()['items']
    assert len(assets) == 2
    serials = {a['serial_number'] for a in assets}
    assert serials == {'SN-CSV-001', 'SN-CSV-002'}

    # Verify audit event was logged
    audit_events = db_session.query(AuditEvent).filter_by(
        company_id=UUID(company['id']),
        action='asset.bulk_import',
    ).all()
    assert len(audit_events) == 1
    assert audit_events[0].changes['imported_count'] == 2


def test_xlsx_bulk_import_preview_and_confirm(client):
    company = create_company(client, 'BULK-XLSX', prefix='XLS')
    setup_company_infrastructure(client, company['id'])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        'category_code', 'serial_number', 'manufacturer', 'model', 'status',
        'site_code', 'location_code', 'department_code', 'cost_center_code',
        'acquisition_date', 'acquisition_cost',
    ])
    ws.append([
        'SERVER', 'SN-XLS-881', 'HP', 'ProLiant DL380', 'in_stock',
        'HQ', 'ROOM-1', 'IT', 'CC-IT',
        '2026-03-01', 5400.00,
    ])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    # Step 1: Preview
    preview_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('servers.xlsx', xlsx_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
    )
    assert preview_res.status_code == 200, preview_res.text
    assert preview_res.json()['valid'] is True
    assert preview_res.json()['valid_rows'] == 1
    token = preview_res.json()['preview_token']

    # Step 2: Confirm
    confirm_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/confirm",
        json={'preview_token': token},
    )
    assert confirm_res.status_code == 200, confirm_res.text
    assert confirm_res.json()['imported_count'] == 1


def test_macro_enabled_spreadsheet_rejected(client):
    company = create_company(client, 'MACRO-REJECT')
    response = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('payload.xlsm', b'PK\x03\x04fake', 'application/vnd.ms-excel.sheet.macroEnabled.12')},
    )
    assert response.status_code == 400
    assert 'Macro-enabled spreadsheets (.xlsm) are not accepted' in response.json()['detail']


def test_validation_missing_required_category_header(client):
    company = create_company(client, 'NO-CAT-HDR')
    csv_text = "serial_number,manufacturer,model\nSN-1,Dell,XPS\n"
    response = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('bad_header.csv', csv_text.encode('utf-8'), 'text/csv')},
    )
    assert response.status_code == 200
    data = response.json()
    assert data['valid'] is False
    assert data['preview_token'] is None
    assert any("Missing required header column 'category_code'" in err['message'] for err in data['errors'])


def test_validation_row_level_diagnostics(client):
    company = create_company(client, 'DIAG')
    setup_company_infrastructure(client, company['id'])

    csv_data = [
        ['category_code', 'serial_number', 'site_code', 'location_code', 'acquisition_cost', 'warranty_start', 'warranty_end', 'status'],
        # Row 2: Inactive / non-existent category
        ['NON_EXISTENT_CAT', 'SN-DIAG-1', 'HQ', 'ROOM-1', '100', '2026-01-01', '2026-02-01', 'in_stock'],
        # Row 3: Location without site
        ['LAPTOP', 'SN-DIAG-2', '', 'ROOM-1', '100', '2026-01-01', '2026-02-01', 'in_stock'],
        # Row 4: Warranty end before start
        ['LAPTOP', 'SN-DIAG-3', 'HQ', 'ROOM-1', '100', '2026-05-01', '2026-01-01', 'in_stock'],
        # Row 5: Negative cost
        ['LAPTOP', 'SN-DIAG-4', 'HQ', 'ROOM-1', '-50.00', '2026-01-01', '2026-02-01', 'in_stock'],
        # Row 6: Invalid status
        ['LAPTOP', 'SN-DIAG-5', 'HQ', 'ROOM-1', '100', '2026-01-01', '2026-02-01', 'flying'],
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows(csv_data)
    csv_bytes = output.getvalue().encode('utf-8')

    response = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('diagnostics.csv', csv_bytes, 'text/csv')},
    )
    assert response.status_code == 200
    data = response.json()
    assert data['valid'] is False
    assert data['valid_rows'] == 0
    assert data['invalid_rows'] == 5
    assert data['preview_token'] is None

    err_map = {err['row_number']: err for err in data['errors']}
    assert 2 in err_map and "Asset category 'NON_EXISTENT_CAT' not found" in err_map[2]['message']
    assert 3 in err_map and "A valid site is required when specifying a location" in err_map[3]['message']
    assert 4 in err_map and "Warranty end date must be on or after warranty start date" in err_map[4]['message']
    assert 5 in err_map and "Acquisition cost must be non-negative" in err_map[5]['message']
    assert 6 in err_map and "Invalid status 'flying'" in err_map[6]['message']


def test_serial_number_batch_and_db_duplicates(client):
    company = create_company(client, 'DUP-SER')
    categories = client.get(f"/api/v1/companies/{company['id']}/categories").json()
    laptop_cat = next(c for c in categories if c['code'] == 'LAPTOP')

    # Pre-create asset in DB with serial SN-EXISTING
    client.post(
        f"/api/v1/companies/{company['id']}/assets",
        json={'category_id': laptop_cat['id'], 'serial_number': 'SN-EXISTING'},
    )

    csv_data = [
        ['category_code', 'serial_number', 'model'],
        ['LAPTOP', 'SN-EXISTING', 'Collides with DB'],
        ['LAPTOP', 'SN-BATCH-DUP', 'First occurrence'],
        ['LAPTOP', 'SN-BATCH-DUP', 'Collides with batch'],
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows(csv_data)

    response = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('dups.csv', output.getvalue().encode('utf-8'), 'text/csv')},
    )
    assert response.status_code == 200
    data = response.json()
    assert data['valid'] is False
    assert any("already exists for an active asset in this company" in e['message'] for e in data['errors'])
    assert any("Duplicate serial number 'SN-BATCH-DUP' within this spreadsheet batch" in e['message'] for e in data['errors'])


def test_tenant_isolation_in_bulk_import(client):
    company_a = create_company(client, 'ISOL-A')
    company_b = create_company(client, 'ISOL-B')

    # Company A creates a custom site HQ-A
    site_a = client.post(
        f"/api/v1/companies/{company_a['id']}/sites",
        json={'code': 'HQ-A', 'name': 'Company A HQ'},
    ).json()

    # Company B tries to import an asset referencing Company A's site HQ-A
    csv_data = [
        ['category_code', 'serial_number', 'site_code'],
        ['LAPTOP', 'SN-CROSS-TENANT', 'HQ-A'],
    ]
    output = io.StringIO()
    csv.writer(output).writerows(csv_data)

    preview_b = client.post(
        f"/api/v1/companies/{company_b['id']}/assets/import/preview",
        files={'file': ('cross.csv', output.getvalue().encode('utf-8'), 'text/csv')},
    )
    assert preview_b.status_code == 200
    data = preview_b.json()
    assert data['valid'] is False
    assert any("Site 'HQ-A' not found or is inactive in company" in e['message'] for e in data['errors'])

    # Company A uploads valid file and gets a preview_token
    csv_data_a = [
        ['category_code', 'serial_number', 'site_code'],
        ['LAPTOP', 'SN-VALID-A', 'HQ-A'],
    ]
    output_a = io.StringIO()
    csv.writer(output_a).writerows(csv_data_a)
    preview_a = client.post(
        f"/api/v1/companies/{company_a['id']}/assets/import/preview",
        files={'file': ('valid_a.csv', output_a.getvalue().encode('utf-8'), 'text/csv')},
    )
    assert preview_a.status_code == 200
    token_a = preview_a.json()['preview_token']
    assert token_a is not None

    # Company B attempts to confirm Company A's token -> 403 Forbidden
    cross_confirm = client.post(
        f"/api/v1/companies/{company_b['id']}/assets/import/confirm",
        json={'preview_token': token_a},
    )
    assert cross_confirm.status_code == 403
    assert 'Preview token does not belong to this company' in cross_confirm.json()['detail']


def test_confirm_with_invalid_or_expired_token(client):
    company = create_company(client, 'BAD-TOKEN')
    response = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/confirm",
        json={'preview_token': '00000000-0000-0000-0000-000000000000'},
    )
    assert response.status_code == 400
    assert 'Invalid or expired import preview token' in response.json()['detail']

def test_formula_sanitization_leading_whitespace():
    from app.asset_import_service import _sanitize_formula

    assert _sanitize_formula('=1+1') == "'=1+1"
    assert _sanitize_formula(' =1+1') == "' =1+1"
    assert _sanitize_formula('\t=CMD()') == "'\t=CMD()"
    assert _sanitize_formula('   @SUM(A1:A10)') == "'   @SUM(A1:A10)"
    assert _sanitize_formula('-5.50') == "'-5.50"
    assert _sanitize_formula('+44123456789') == "'+44123456789"
    assert _sanitize_formula('Normal Text') == 'Normal Text'
    assert _sanitize_formula(123) == '123'
    assert _sanitize_formula(None) == ''


def test_bounded_file_read_oversized_file(client):
    company = create_company(client, 'OVERSIZE')
    from app.asset_import_service import MAX_FILE_SIZE_BYTES

    oversized_bytes = b'a' * (MAX_FILE_SIZE_BYTES + 50)
    response = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('big.csv', oversized_bytes, 'text/csv')},
    )
    assert response.status_code == 413
    assert 'File size exceeds maximum allowed limit' in response.json()['detail']


def test_concurrent_confirm_duplicate_protection(client):
    company = create_company(client, 'RACE-LOCK')
    csv_data = [
        ['category_code', 'serial_number'],
        ['LAPTOP', 'SN-RACE-LOCK-1'],
    ]
    output = io.StringIO()
    csv.writer(output).writerows(csv_data)
    preview_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('race.csv', output.getvalue().encode('utf-8'), 'text/csv')},
    )
    assert preview_res.status_code == 200
    token = preview_res.json()['preview_token']

    # Simulate an active in-flight request by manually setting is_processing = True
    from app.asset_import_service import _STAGING_LOCK, _STAGING_STORE
    with _STAGING_LOCK:
        assert token in _STAGING_STORE
        _STAGING_STORE[token].is_processing = True

    # Concurrent confirm attempt must be rejected with 409
    concurrent_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/confirm",
        json={'preview_token': token},
    )
    assert concurrent_res.status_code == 409
    assert 'already being processed' in concurrent_res.json()['detail']

    # Once lock/processing is released, confirm can proceed
    with _STAGING_LOCK:
        _STAGING_STORE[token].is_processing = False

    confirm_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/confirm",
        json={'preview_token': token},
    )
    assert confirm_res.status_code == 200
    assert confirm_res.json()['imported_count'] == 1


def test_confirm_revalidates_serial_numbers_emerging_after_preview(client):
    company = create_company(client, 'RACE-SER')
    categories = client.get(f"/api/v1/companies/{company['id']}/categories").json()
    laptop_cat = next(c for c in categories if c['code'] == 'LAPTOP')

    csv_data = [
        ['category_code', 'serial_number'],
        ['LAPTOP', 'SN-INTERIM-COLLIDE'],
    ]
    output = io.StringIO()
    csv.writer(output).writerows(csv_data)

    preview_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/preview",
        files={'file': ('race_ser.csv', output.getvalue().encode('utf-8'), 'text/csv')},
    )
    assert preview_res.status_code == 200
    token = preview_res.json()['preview_token']

    # Now create an asset in the company with that serial BEFORE confirmation
    client.post(
        f"/api/v1/companies/{company['id']}/assets",
        json={'category_id': laptop_cat['id'], 'serial_number': 'SN-INTERIM-COLLIDE'},
    )

    # Confirmation must catch the collision and abort with 409
    confirm_res = client.post(
        f"/api/v1/companies/{company['id']}/assets/import/confirm",
        json={'preview_token': token},
    )
    assert confirm_res.status_code == 409
    assert 'Serial numbers conflicted with records created after preview' in confirm_res.json()['detail']


def test_staging_store_memory_ceiling_eviction():
    import time
    from uuid import uuid4
    from app.asset_import_service import (
        MAX_STAGED_IMPORTS,
        StagedImport,
        _clean_staging_store,
        _STAGING_LOCK,
        _STAGING_STORE,
    )

    with _STAGING_LOCK:
        _STAGING_STORE.clear()
        now = time.time()
        # Seed MAX_STAGED_IMPORTS + 10 items
        for i in range(MAX_STAGED_IMPORTS + 10):
            token = f"test-token-{i}"
            _STAGING_STORE[token] = StagedImport(
                token=token,
                company_id=uuid4(),
                actor_id=uuid4(),
                created_at=now + i,
                rows=[],
                row_count=0,
                is_processing=False,
            )
        assert len(_STAGING_STORE) == MAX_STAGED_IMPORTS + 10
        _clean_staging_store()
        assert len(_STAGING_STORE) <= MAX_STAGED_IMPORTS
        # Oldest tokens 0 through 9 should have been evicted
        assert 'test-token-0' not in _STAGING_STORE
        assert 'test-token-9' not in _STAGING_STORE
        assert f"test-token-{MAX_STAGED_IMPORTS + 9}" in _STAGING_STORE
        _STAGING_STORE.clear()
